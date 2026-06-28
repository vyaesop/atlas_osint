"""Turn an uploaded document into plain text.

TXT / CSV / JSON are handled with the stdlib (no dependencies). PDF and DOCX
need optional libraries (``pypdf`` / ``python-docx``) which are imported lazily,
so the rest of ingestion works even if they aren't installed.
"""
from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass

SUPPORTED_EXTENSIONS = {"txt", "csv", "json", "pdf", "docx"}

# Content-type → canonical kind, for when the filename has no useful extension.
_CONTENT_TYPE_KIND = {
    "text/plain": "txt",
    "text/csv": "csv",
    "application/json": "json",
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
}


class UnsupportedDocumentError(ValueError):
    pass


@dataclass(slots=True)
class ParsedDocument:
    text: str
    kind: str
    char_count: int


def _kind_for(filename: str | None, content_type: str | None) -> str:
    if filename and "." in filename:
        ext = filename.rsplit(".", 1)[-1].lower()
        if ext in SUPPORTED_EXTENSIONS:
            return ext
    if content_type and content_type in _CONTENT_TYPE_KIND:
        return _CONTENT_TYPE_KIND[content_type]
    raise UnsupportedDocumentError(
        f"Unsupported document (filename={filename!r}, content_type={content_type!r}). "
        f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}."
    )


def _parse_csv(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))
    # Flatten to readable lines so the extractor sees natural-ish text.
    return "\n".join(", ".join(cell.strip() for cell in row) for row in rows if any(row))


def _parse_json(data: bytes) -> str:
    obj = json.loads(data.decode("utf-8", errors="replace"))

    def walk(node, lines: list[str]) -> None:
        if isinstance(node, dict):
            for k, v in node.items():
                if isinstance(v, (dict, list)):
                    walk(v, lines)
                else:
                    lines.append(f"{k}: {v}")
        elif isinstance(node, list):
            for item in node:
                walk(item, lines)
        else:
            lines.append(str(node))

    lines: list[str] = []
    walk(obj, lines)
    return "\n".join(lines)


def _parse_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader  # lazy import
    except ImportError as exc:  # pragma: no cover
        raise UnsupportedDocumentError(
            "PDF parsing requires the 'pypdf' package."
        ) from exc
    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _parse_docx(data: bytes) -> str:
    try:
        import docx  # python-docx, lazy import
    except ImportError as exc:  # pragma: no cover
        raise UnsupportedDocumentError(
            "DOCX parsing requires the 'python-docx' package."
        ) from exc
    document = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs)


def parse_document(
    data: bytes, *, filename: str | None = None, content_type: str | None = None
) -> ParsedDocument:
    kind = _kind_for(filename, content_type)
    if kind == "txt":
        text = data.decode("utf-8", errors="replace")
    elif kind == "csv":
        text = _parse_csv(data)
    elif kind == "json":
        text = _parse_json(data)
    elif kind == "pdf":
        text = _parse_pdf(data)
    elif kind == "docx":
        text = _parse_docx(data)
    else:  # pragma: no cover - guarded by _kind_for
        raise UnsupportedDocumentError(kind)

    text = text.strip()
    return ParsedDocument(text=text, kind=kind, char_count=len(text))
